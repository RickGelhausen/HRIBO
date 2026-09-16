"""The Excel export mirrors the five visual cross-condition report views."""

from types import SimpleNamespace

from openpyxl import load_workbook

from generate_diffex_summary import write_excel


def _detection(state, mean_cpm=10):
    return {
        "state": state, "replicates": 2, "passing_replicates": 2,
        "mean_count": 20, "mean_cpm": mean_cpm,
        "normalized_count": mean_cpm, "replicate_fraction": 1,
        "samples": [],
    }


def _effect(state, log2fc=None, padj=None):
    return {"state": state, "log2fc": log2fc, "padj": padj}


def _fixture():
    feature_id = "chr:1-90:+"
    rows = [{
        "feature_id": feature_id,
        "conditions": {
            "WT": {"RNA": _detection("detected"), "RIBO": _detection("detected")},
            "Mut": {"RNA": _detection("not_detected", 0), "RIBO": _detection("uncertain", 2)},
        },
        "contrasts": {
            "Mut-WT": {
                "RNA": _effect("up", 2.25, 0.01),
                "RIBO": _effect("down", -1.75, 0.02),
                "TE": _effect("not_significant", 0.4, 0.8),
            },
        },
    }]
    coordinates = {feature_id: {
        "seqid": "chr", "start": 1, "end": 90, "strand": "+",
        "type": "CDS", "name": "geneA", "locus_tag": "b0001",
    }}
    settings = SimpleNamespace(
        min_count=10, min_cpm=1, min_replicates=2,
        padj_cutoff=0.05, log2fc_cutoff=1,
    )
    return rows, coordinates, settings


def test_workbook_mirrors_each_html_view_and_is_sortable(tmp_path):
    rows, coordinates, settings = _fixture()
    path = tmp_path / "condition_overview.xlsx"

    write_excel(path, rows, coordinates, ["WT", "Mut"], ["Mut-WT"], settings)

    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == [
        "RNA_detection", "RIBO_detection", "RNA_change", "RIBO_change",
        "TE_change", "README",
    ]
    detection = workbook["RNA_detection"]
    assert [cell.value for cell in detection[1]] == [
        "Locus tag", "Identifier", "WT", "Mut",
    ]
    for sheet_name in workbook.sheetnames[:-1]:
        sheet = workbook[sheet_name]
        assert [sheet.cell(1, column).value for column in (1, 2)] == [
            "Locus tag", "Identifier",
        ]
        assert sheet.freeze_panes == "C2"
    assert detection["A2"].value == "b0001"
    assert detection["B2"].value == "chr:1-90:+"
    assert detection["C2"].value == "Detected"
    assert detection["D2"].value == "Not detected"
    assert detection.freeze_panes == "C2"
    assert detection.auto_filter.ref == "A1:D2"

    rna_change = workbook["RNA_change"]
    ribo_change = workbook["RIBO_change"]
    assert rna_change["C2"].value == 2.25
    assert ribo_change["C2"].value == -1.75
    assert "↑" in rna_change["C2"].number_format
    assert "↓" in ribo_change["C2"].number_format
    assert workbook["TE_change"]["C2"].value == "No directional call"
    assert rna_change["C2"].fill.fgColor.rgb.endswith("F7C8C4")
    assert ribo_change["C2"].fill.fgColor.rgb.endswith("C7DCFA")


def test_excel_writes_identity_text_literally(tmp_path):
    rows, coordinates, settings = _fixture()
    feature_id = rows[0]["feature_id"]
    malicious_identifier = '=HYPERLINK("identifier")'
    coordinates[feature_id]["locus_tag"] = '=HYPERLINK("locus")'
    coordinates[malicious_identifier] = coordinates.pop(feature_id)
    rows[0]["feature_id"] = malicious_identifier
    path = tmp_path / "condition_overview.xlsx"

    write_excel(path, rows, coordinates, ["WT", "Mut"], ["Mut-WT"], settings)

    sheet = load_workbook(path, data_only=False)["RNA_detection"]
    assert sheet["A2"].value == '=HYPERLINK("locus")'
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == malicious_identifier
    assert sheet["B2"].data_type == "s"
